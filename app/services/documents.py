"""Распознавание документов клиента: оборотно-сальдовая ведомость по счёту 62 (выгрузка из 1С, XLSX или CSV).

Без OCR и без ИИ: файл из 1С уже структурирован, правила здесь надёжнее модели. Находим обороты по дебету
счёта 62 за период (отгрузки покупателям) и сверяем с выручкой из заявки — с НДС 20 % и без, потому что
выручка в отчётности указывается без НДС, а отгрузки в ОСВ — с ним. Результат видит только аналитик:
это подсказка, решение остаётся за ним. Сканы и PDF — в «Потенциале развития».
"""
from __future__ import annotations

import csv
import os
import re
from dataclasses import dataclass
from functools import lru_cache
from typing import Any

OSV_62 = "Оборотно-сальдовая ведомость по счёту 62"
TOLERANCE = 0.10  # расхождение до 10 % считаем совпадением: разные периоды, округления, авансы
VAT = 1.2
_PERIOD = re.compile(r"(\d{2}\.\d{2}\.\d{4})\s*[-–—]\s*(\d{2}\.\d{2}\.\d{4})")


class NotRecognized(Exception):
    pass


@dataclass(frozen=True)
class OsvCheck:
    recognized: bool
    period: str | None = None
    debit: int | None = None  # обороты по дебету 62 за период
    declared: int | None = None  # выручка из заявки
    diff_pct: int | None = None  # наименьшее расхождение: с НДС или без
    with_vat: bool | None = None  # сошлось с учётом НДС
    problem: str | None = None

    @property
    def verdict(self) -> str:
        if not self.recognized or self.diff_pct is None:
            return "unknown"
        return "match" if self.diff_pct <= TOLERANCE * 100 else "mismatch"


def _cell_num(v: Any) -> float | None:
    if isinstance(v, (int, float)):
        return float(v)
    if not isinstance(v, str):
        return None
    s = v.replace(" ", "").replace(" ", "").replace(",", ".").strip()
    try:
        return float(s) if s not in ("", "-", "—") else None
    except ValueError:
        return None


def _text(v: Any) -> str:
    return str(v).strip().lower().replace("ё", "е") if v is not None else ""


def read_rows(path: str, filename: str) -> list[list[Any]]:
    ext = os.path.splitext(filename)[1].lower()
    if ext == ".xlsx":
        from openpyxl import load_workbook  # тяжёлый импорт — только когда нужен

        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            return [list(r) for r in wb.active.iter_rows(values_only=True)]
        finally:
            wb.close()
    if ext == ".csv":
        raw = open(path, "rb").read()
        for enc in ("utf-8-sig", "cp1251"):
            try:
                text = raw.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        else:
            raise NotRecognized("неизвестная кодировка CSV")
        delim = ";" if text.count(";") >= text.count(",") else ","
        return [row for row in csv.reader(text.splitlines(), delimiter=delim)]
    raise NotRecognized("распознаём только XLSX и CSV из 1С")


def parse_osv62(rows: list[list[Any]]) -> tuple[str | None, int]:
    """Период и обороты по дебету 62 из ОСВ в формате 1С."""
    period = None
    for row in rows[:12]:
        for v in row:
            m = _PERIOD.search(str(v or ""))
            if m:
                period = f"{m.group(1)} – {m.group(2)}"
                break
        if period:
            break
    head = next(((r, c) for r, row in enumerate(rows) for c, v in enumerate(row)
                 if "оборот" in _text(v) and "период" in _text(v)), None)
    if head is None:
        raise NotRecognized("не нашли колонку «Обороты за период»")
    hr, hc = head
    debit_col = None
    for row in rows[hr:hr + 2]:
        debit_col = next((c for c, v in enumerate(row) if c >= hc and _text(v) == "дебет"), None)
        if debit_col is not None:
            break
    if debit_col is None:
        raise NotRecognized("не нашли «Дебет» под оборотами")
    body = rows[hr + 1:]
    total = next((row for row in body if row and _text(row[0]).startswith("итого")), None) \
        or next((row for row in body if row and _text(row[0]) == "62"), None)
    if total is None:
        raise NotRecognized("не нашли строку «Итого» или «62»")
    debit = _cell_num(total[debit_col]) if debit_col < len(total) else None
    if debit is None:
        raise NotRecognized("в строке итогов нет оборота по дебету")
    return period, int(round(debit))


def compare(debit: int, declared: int | None) -> tuple[int | None, bool | None]:
    if not declared:
        return None, None
    with_vat = abs(debit / VAT - declared) / declared
    without = abs(debit - declared) / declared
    return int(round(min(with_vat, without) * 100)), with_vat <= without


@lru_cache(maxsize=256)
def _check_cached(path: str, filename: str, sha256: str, declared: int | None) -> OsvCheck:
    try:
        period, debit = parse_osv62(read_rows(path, filename))
    except NotRecognized as e:
        return OsvCheck(False, problem=str(e))
    except Exception:  # noqa: BLE001 — битый файл не должен ронять карточку
        return OsvCheck(False, problem="файл не читается")
    diff, with_vat = compare(debit, declared)
    return OsvCheck(True, period=period, debit=debit, declared=declared, diff_pct=diff, with_vat=with_vat)


def check(kind: str, filename: str, path: str, sha256: str, declared_revenue: int | None) -> OsvCheck | None:
    """Проверка документа для карточки аналитика. None — этот документ не распознаём."""
    if kind != OSV_62 or not os.path.exists(path):
        return None
    return _check_cached(path, filename, sha256, declared_revenue)
